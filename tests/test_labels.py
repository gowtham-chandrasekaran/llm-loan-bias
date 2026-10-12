import json, os, sys
import numpy as np, pandas as pd, pytest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
from llmbias import labels as L

def test_known_kappa_ac1_pabak():
    a = ["1"] * 20 + ["0"] * 80; b = ["1"] * 15 + ["0"] * 5 + ["1"] * 5 + ["0"] * 75
    # po=0.9; pe=0.2*0.2+0.8*0.8=0.68 -> kappa=0.6875
    assert abs(L.cohen_kappa(a, b) - 0.6875) < 1e-9
    # AC1: pi=0.2, pe=2*0.2*0.8/1=0.32 -> (0.9-0.32)/0.68
    assert abs(L.gwet_ac1(a, b) - (0.9 - 0.32) / 0.68) < 1e-9
    assert abs(L.pabak(a, b) - 0.8) < 1e-9
    assert L.cohen_kappa(["1", "0"], ["1", "0"]) == 1.0
    lo, hi = L.boot_ci(a, b, L.cohen_kappa, B=500, seed=1); assert lo < 0.6875 < hi
    assert L.boot_ci(a, b, L.cohen_kappa, B=500, seed=1) == (lo, hi)   # seeded

def test_norm_label_rejects():
    assert L.norm_label(1) == "1" and L.norm_label(0.0) == "0" and L.norm_label(" Unsure ") == "unsure"
    for bad in ("yes", "2", "maybe", "0.5", -1):
        with pytest.raises(L.LabelError): L.norm_label(bad)

def _fake_run(tmp_path):
    rd = tmp_path / "RUNX"; rd.mkdir(); calls, judg = [], []
    rng = np.random.default_rng(0)
    for i in range(300):
        v = ["base", "name_b", "zip_b", "anchor"][i % 4]; cid = f"c{i}"
        calls.append({"call_id": cid, "setup": 3, "variant": v, "parsed": {"decision": "deny"}})
        judg.append({"item_id": f"i{i:04d}", "call_id": cid, "letter_redacted": f"Dear [NAME], letter {i}", "hedging": bool(rng.random() < .3),
                     "hedging_status": "done", "spans": [{"span_redacted": f"span {i}", "mentions_proxy": bool(rng.random() < .5), "status": "done"}]})
    (rd / "calls.jsonl").write_text("\n".join(map(json.dumps, calls))); (rd / "judgments.jsonl").write_text("\n".join(map(json.dumps, judg)))
    return str(rd)

@pytest.fixture
def env(tmp_path, monkeypatch):
    priv = tmp_path / "private"; monkeypatch.setattr(L, "PRIVATE", str(priv))
    return _fake_run(tmp_path), str(tmp_path / "out"), priv

def test_export_roundtrip_and_validation(env):
    from openpyxl import load_workbook
    rd, out, priv = env
    paths, kp = L.export(rd, out)
    assert oct(os.stat(kp).st_mode)[-3:] == "600" and oct(os.stat(os.path.dirname(kp)).st_mode)[-3:] == "700"
    ax, bx = paths["A"][0], paths["B"][0]
    assert paths["A"][2] == 180 and paths["A"][3] == {"hedging": 120, "proxy_mention": 60}
    wa, wb_ = load_workbook(ax), load_workbook(bx)
    ws = wa["Items"]; assert [c.value for c in ws[1]] == ["item_id", "task", "text", "label", "comment"]
    dv = ws.data_validations.dataValidation[0]
    assert dv.formula1 == '"1,0,unsure"' and dv.showErrorMessage and dv.errorStyle == "stop" and "D2:D181" in str(dv.sqref)
    assert {"Rubric", "About"} <= set(wa.sheetnames) and L.RUBRIC_SHA_PREFIX in wa["Rubric"]["A1"].value
    ids_a = [r[0].value for r in ws.iter_rows(min_row=2)]; ids_b = [r[0].value for r in wb_["Items"].iter_rows(min_row=2)]
    assert not set(ids_a) & set(ids_b)
    key = json.load(open(kp))
    ord_a = [key["labelers"]["A"]["map"][i]["internal_id"] for i in ids_a]; ord_b = [key["labelers"]["B"]["map"][i]["internal_id"] for i in ids_b]
    assert sorted(ord_a) == sorted(ord_b) and ord_a != ord_b
    L.check_no_leak(out)
    # fill: A = judge, B = judge with some flips and unsure; round trip via xlsx
    for wbk, lab in ((wa, "A"), (wb_, "B")):
        ws = wbk["Items"]
        for k, row in enumerate(ws.iter_rows(min_row=2)):
            j = key["labelers"][lab]["map"][row[0].value]["judge"]
            val = j if lab == "A" else (1 - j if k % 10 == 0 else j)
            if lab == "B" and k == 3: val = "unsure"; row[4].value = None
            row[3].value = val
        wbk.save(ax if lab == "A" else bx)
    res = L.agreement(rd, ax, bx, B=300)
    assert len(res["flags_unsure_without_comment"]["B"]) == 1
    h = res["per_task"]["hedging"]; assert h["A_vs_judge"]["kappa"] == 1.0 and h["human_human"]["n"] == 120
    assert os.path.exists(res["adjudication_file"])
    # adjudicate: resolve all but one to the A label
    from openpyxl import load_workbook as lw
    adj = lw(res["adjudication_file"]); ws = adj["Disagreements"]
    for k, row in enumerate(ws.iter_rows(min_row=2)):
        if k > 0: row[7].value = row[3].value if row[3].value != "unsure" else row[5].value
    adj.save(res["adjudication_file"])
    res2 = L.agreement(rd, ax, bx, adjudication_path=res["adjudication_file"], adj_out=str(env[1]) + "/adj2.xlsx", B=300)
    tot = sum(v["unresolved_or_both_unsure"] for v in res2["consensus"].values()); assert tot == 1

def test_both_unsure_excluded_one_unsure_disagreement(env):
    rd, out, priv = env
    paths, kp = L.export(rd, out); key = json.load(open(kp))
    def write(lab, fn):
        df = pd.read_csv(paths[lab][1], dtype=str, keep_default_na=False)
        m = key["labelers"][lab]["map"]
        df["label"] = [fn(m[i]) for i in df.item_id]; df["comment"] = ["x" if l == "unsure" else "" for l in df.label]
        df.to_csv(paths[lab][1], index=False)
    target = sorted(key["selected_internal_ids"])[:2]   # t0 both unsure, t1 only A unsure
    write("A", lambda r: "unsure" if r["internal_id"] in target else str(r["judge"]))
    write("B", lambda r: "unsure" if r["internal_id"] == target[0] else str(r["judge"]))
    res = L.agreement(rd, paths["A"][1], paths["B"][1], B=200)
    allt = res["per_task"]
    assert sum(t["both_unsure_excluded"] for t in allt.values()) == 1
    assert sum(res["n_disagreements"].values()) == 1

def test_csv_invalid_value_rejected(env):
    rd, out, priv = env
    paths, _ = L.export(rd, out)
    df = pd.read_csv(paths["A"][1], dtype=str, keep_default_na=False); df["label"] = "1"; df.loc[5, "label"] = "yes"
    df.to_csv(paths["A"][1], index=False)
    with pytest.raises(L.LabelError): L.read_labels(paths["A"][1])

def test_calibration_outside_180(env):
    rd, out, priv = env
    L.export(rd, out)
    p = L.draw_calibration(rd, outdir=os.path.join(out, "cal"))
    k = json.load(open(os.path.join(priv, "RUNX", "labeling_pilot-v1_calibration_key.json")))
    used = set(json.load(open(os.path.join(priv, "RUNX", "labeling_pilot-v1_key.json")))["selected_internal_ids"])
    ids = {v["internal_id"] for v in k["map"].values()}
    assert len(ids) == 20 and not ids & used and k["seed"] == 15803
