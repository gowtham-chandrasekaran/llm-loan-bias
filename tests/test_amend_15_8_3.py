import json, math, os, stat
import numpy as np, pandas as pd, pytest
import llmbias.labels as LB
from llmbias.blind import is_salutation_only, normalize_span, filter_spans, has_substantive_placeholder, detect_proxy_spans
from llmbias.prompts import judge_mention

def test_salutation_only_regex():
    for s in ["[NAME],", "[NAME]", "Dear [NAME],", "Dear Mr. [NAME],", "Ms. [NAME]", "Dear Mrs. [NAME]"]: assert is_salutation_only(s)
    for s in ["Dear [NAME], your loan in [ZIP] is approved.", "Applicant [NAME] has a DTI of 45%.", "[ZIP],", "Dear Dr. [NAME],"]: assert not is_salutation_only(s)

def test_filter_and_dedup():
    sp = ["[NAME],", "Your property in [ZIP] qualifies.", "your property in [ZIP]  qualifies", "Dear Ms. [NAME],", "Other text."]
    kept, ns, nd = filter_spans(sp)
    assert kept == ["Your property in [ZIP] qualifies.", "Other text."] and ns == 2 and nd == 1
    assert normalize_span("  A  b. ") == "a b"
    assert has_substantive_placeholder("Located in [STATE].") and not has_substantive_placeholder("[NAME],") and not has_substantive_placeholder("No cue.")

def test_detect_proxy_spans_drops_name_line():
    t = "Latoya Harris,\nYour home in ZIP 19104 meets standards. Your home in ZIP 19104 meets standards."
    out = detect_proxy_spans(t, "Latoya Harris", "19104")
    assert [s["span_redacted"] for s in out] == ["Your home in ZIP [ZIP] meets standards."]

def test_judge_prompt_b2():
    _, u = judge_mention("The property in [ZIP] meets our standards.")
    for k in ("[ZIP]", "[STATE]", "[CITY]", "[NAME]", "[ADDRESS]", "[NEIGHBORHOOD]", "MENTIONS", "ECOA", "mentions_proxy"): assert k in u

def test_kappa_undefined_when_pe_one():
    assert math.isnan(LB.cohen_kappa(["0"] * 10, ["0"] * 10))
    assert math.isnan(LB.cohen_kappa([], []))
    assert LB.cohen_kappa(["1", "0"], ["1", "0"]) == 1.0

def test_gate_not_evaluable():
    b = LB.agree_block(["0"] * 50, ["0"] * 50, B=50)
    assert b["gate_status"] == LB.NOT_EVALUABLE and b["gate_0.70_pass"] is False and b["kappa"] is None
    a = ["1"] * 4 + ["0"] * 96; b = LB.agree_block(a, list(a), B=50)        # kappa = 1 but < 5 positives
    assert b["kappa"] == 1.0 and b["gate_status"] == LB.NOT_EVALUABLE and not b["gate_0.70_pass"]
    assert math.isnan(b["ac1"]) and math.isnan(b["pabak"]) and "4" in b["gate_reason"]
    a = ["1"] * 30 + ["0"] * 70; b = LB.agree_block(a, list(a), B=50)
    assert b["gate_status"] == "pass" and b["gate_0.70_pass"] and not math.isnan(b["ac1"])
    a2 = ["1"] * 15 + ["0"] * 15 + ["1"] * 15 + ["0"] * 55; b = LB.agree_block(a, a2, B=50)
    assert b["gate_status"] == "fail"

def _fake_pool():
    rows = []; rng = np.random.default_rng(0)
    for g in ("base", "name", "zip"):
        for i in range(120):
            t = f"Letter {g} {i}." + (" Unfortunately we cannot." if i % 10 == 0 else "")
            rows.append(dict(internal_id=f"h{g}{i}", task="hedging", variant=g, call_id=f"c{g}{i}", text=t, judge=int(i < 7)))
        for i in range(120):
            t = ["[NAME],", f"Your home in [ZIP] block {g}{i} qualifies.", f"Income reason {g}{i} applies."][i % 3]
            rows.append(dict(internal_id=f"h{g}{i}#s0", task="proxy_mention", variant=g, call_id=f"c{g}{i}", text=t, judge=int(i == 1)))
    return pd.DataFrame(rows)

def test_select_v2_strata():
    items, log = LB.select_v2(_fake_pool(), seed=1)
    h = items[items.task == "hedging"]; s = items[items.task == "proxy_mention"]
    assert len(h) == 120 and len(s) == 100
    assert (h.stratum == "m1b:judge_positive").sum() == 21 and (h.stratum == "m1b:borderline_cue").sum() == 30
    assert not s.text.map(is_salutation_only).any() and s.text.map(normalize_span).is_unique
    assert s.stratum.str.startswith("m4:cue").sum() == 50 and s.stratum.str.startswith("m4:no_cue").sum() == 50
    assert set(s.variant) == {"base", "name", "zip"} and items.internal_id.is_unique

def test_export_v2_private_and_no_leak(tmp_path, monkeypatch):
    monkeypatch.setattr(LB, "PRIVATE", str(tmp_path / "priv")); monkeypatch.setattr(LB, "item_pool", lambda run_dir: _fake_pool())
    out = tmp_path / "lab"; paths, kp, sp, log = LB.export_v2(str(tmp_path / "runX"), outdir=str(out))
    for p in (kp, sp): assert stat.S_IMODE(os.stat(p).st_mode) == 0o600 and str(out) not in p
    assert sorted(os.listdir(out)) == ["LABELING_METADATA.json", "labeler_A.csv", "labeler_A.xlsx", "labeler_B.csv", "labeler_B.xlsx"]
    a, b = pd.read_csv(out / "labeler_A.csv"), pd.read_csv(out / "labeler_B.csv")
    assert sorted(a.text) == sorted(b.text) and list(a.text) != list(b.text) and not set(a.item_id) & set(b.item_id)
    key = json.load(open(kp)); assert key["labelers"]["A"]["order_seed"] == 15811 and key["labelers"]["B"]["order_seed"] == 15812
    from openpyxl import load_workbook
    wb = load_workbook(out / "labeler_A.xlsx"); ws = wb["Items"]
    assert [c.value for c in ws[1]] == ["item_id", "task", "text", "label", "comment"]
    assert ws.data_validations.dataValidation[0].formula1 == '"1,0,unsure"'
    assert "v1.1" in wb["Rubric"]["A1"].value and "f4747b61" in wb["Rubric"]["A1"].value
    (out / "notes_selection.txt").write_text("x")
    with pytest.raises(AssertionError): LB.check_no_disclosure(str(out))
